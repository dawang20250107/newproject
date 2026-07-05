"""内部往来核对（金蝶云星空补位模块）。

数据源：金蝶「核算维度明细账」（核算维度＝往来单位/内部往来单位）按记账主体
（事业部/总部账簿）导出的 xlsx，逐主体上传进本系统。

核对逻辑（镜像原则）：
- A 主体账上「对 B 的应收」（资产类科目，借增）必然对应 B 主体账上
  「对 A 的应付」（负债类科目，贷增）。
- 每行 signed = 借-贷（债权为正）。A 对 B 的 signed 合计 与 B 对 A 的
  signed 合计 互为相反数即核平；两者之和即差异。
- 两两明细比对：金额相等且方向互镜（signed 相反）的行自动配对标记，
  未配对行即差异来源（单边入账/金额录错/期间错位）。
"""
import datetime
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.views.decorators.csrf import csrf_exempt

from .models import InternalBatch, InternalEntry, BUSINESS_UNITS, VALID_BUSINESS_UNITS
from .views import (
    cw_required, ok, err, _page_denied, _can_upload, _can_delete,
    _can_access_bu, IMPORT_SIZE_LIMIT,
)

# 对方单位原文 → 事业部映射：先精确匹配全名，再按识别词干包含匹配。
# 词干有序：更长/更特异的在前（「多式联运」不含「运输」二字序列，互不误伤）。
_CP_STEMS = [
    ('多式联运', '多式联运事业部'),
    ('供应链', '供应链事业部'),
    ('劳务', '劳务事业部'),
    ('运输', '运输事业部'),
    ('自营', '自营事业部'),
    ('阔展', '阔展事业部'),
    ('总部', '集团总部'),
    ('集团', '集团总部'),
]

# 金蝶明细账里的小计/汇总行摘要（不是业务发生额，须跳过）
_SKIP_SUMMARIES = {'期初余额', '本期合计', '本年累计', '本月合计', '本期发生', '累计', '合计'}


def _map_counterparty(raw):
    """往来单位原文 → 事业部名；识别不了返回 ''（保留原文供人工确认）。"""
    s = (raw or '').strip()
    if not s:
        return ''
    if s in VALID_BUSINESS_UNITS:
        return s
    for stem, bu in _CP_STEMS:
        if stem in s:
            return bu
    return ''


def _detect_internal_ledger(ws):
    """识别金蝶「核算维度明细账（往来单位）」表头。
    必需列：往来单位维度、摘要、借方、贷方；可选：日期、凭证字号、科目编码、科目名称。
    返回 (data_start, col_map) 或 (None, {})。"""
    dim_names = ('往来单位', '内部往来单位', '往来单位名称', '内部单位', '往来组织', '核算维度')
    for ri in range(1, min(10, ws.max_row + 1)):
        cm = {}
        for ci in range(1, min(ws.max_column + 1, 30)):
            v = str(ws.cell(row=ri, column=ci).value or '').strip()
            if not v:
                continue
            if v in dim_names or any(v.startswith(d) for d in dim_names[:5]):
                cm['cp'] = ci
            elif v == '摘要':
                cm['summary'] = ci
            elif v in ('借方', '借方金额'):
                cm['debit'] = ci
            elif v in ('贷方', '贷方金额'):
                cm['credit'] = ci
            elif v == '日期':
                cm['date'] = ci
            elif v in ('凭证字号', '凭证号'):
                cm['voucher'] = ci
            elif v == '科目编码':
                cm['code'] = ci
            elif v == '科目名称':
                cm['name'] = ci
        if all(k in cm for k in ('cp', 'summary', 'debit', 'credit')):
            return ri + 1, cm
    return None, {}


def _cell_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v or '').strip()
    for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y年%m月%d日', '%Y.%m.%d'):
        try:
            return datetime.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _cell_dec(v):
    try:
        s = str(v if v is not None else 0).replace(',', '').strip() or '0'
        return Decimal(s)
    except (InvalidOperation, TypeError):
        return Decimal('0')


@csrf_exempt
@cw_required()
def internal_upload(request):
    """POST multipart：file + bu + year + month。同期间同主体重复上传整体替换。"""
    if request.method != 'POST':
        return err('方法不允许', 405)
    denied = _page_denied(request, 'internal')
    if denied:
        return denied
    if not (request.pk_role == 'super_admin' or _can_upload(request)):
        return err('无上传权限', 403)
    bu = (request.POST.get('bu') or '').strip()
    if bu not in VALID_BUSINESS_UNITS:
        return err('请选择有效的记账主体（事业部/总部）')
    if not _can_access_bu(request, bu):
        return err('您无权为该主体上传数据', 403)
    try:
        year = int(request.POST.get('year', ''))
        month = int(request.POST.get('month', ''))
        assert 2000 <= year <= 2100 and 1 <= month <= 12
    except Exception:
        return err('年份或月份无效')
    f = request.FILES.get('file')
    if not f:
        return err('请上传文件')
    if f.size > IMPORT_SIZE_LIMIT:
        return err('文件过大（上限5MB）')
    try:
        import openpyxl
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
    except Exception:
        return err('文件格式错误，请上传金蝶导出的 Excel(.xlsx)')

    data_start, cm = _detect_internal_ledger(ws)
    if data_start is None:
        return err('无法识别为「核算维度明细账（往来单位）」：表头需含「往来单位」「摘要」「借方」「贷方」列。'
                   '请在金蝶总账 → 核算维度明细账，核算维度选“往来单位”后导出')

    rows, skipped, unmatched = [], 0, {}
    for ri in range(data_start, ws.max_row + 1):
        summ = str(ws.cell(row=ri, column=cm['summary']).value or '').strip()
        if summ in _SKIP_SUMMARIES or '结转' in summ:
            skipped += 1
            continue
        debit = _cell_dec(ws.cell(row=ri, column=cm['debit']).value)
        credit = _cell_dec(ws.cell(row=ri, column=cm['credit']).value)
        if not debit and not credit:
            skipped += 1
            continue
        raw_cp = str(ws.cell(row=ri, column=cm['cp']).value or '').strip()
        # 组合维度串（如「往来单位:XX公司」）取冒号后值
        if ':' in raw_cp or '：' in raw_cp:
            raw_cp = raw_cp.replace('：', ':').split(':')[-1].strip()
        cp = _map_counterparty(raw_cp)
        if cp == bu:
            # 自己对自己：多为维度误选，跳过并计入未识别提示
            unmatched[raw_cp + '（与记账主体相同）'] = unmatched.get(raw_cp + '（与记账主体相同）', 0) + 1
            skipped += 1
            continue
        if not cp:
            unmatched[raw_cp or '（空）'] = unmatched.get(raw_cp or '（空）', 0) + 1
        code = str(ws.cell(row=ri, column=cm['code']).value or '').strip() if 'code' in cm else ''
        side = 'ap' if code.split('.')[0].startswith('2') else 'ar'
        rows.append(InternalEntry(
            business_unit=bu, counterparty=cp, counterparty_raw=raw_cp[:200],
            year=year, month=month,
            biz_date=_cell_date(ws.cell(row=ri, column=cm['date']).value) if 'date' in cm else None,
            voucher_no=str(ws.cell(row=ri, column=cm['voucher']).value or '').strip()[:64] if 'voucher' in cm else '',
            subject_code=code[:32],
            subject_name=str(ws.cell(row=ri, column=cm['name']).value or '').strip()[:100] if 'name' in cm else '',
            summary=summ[:300], debit=debit, credit=credit, side=side,
        ))
    if not rows:
        return err('未解析到任何内部往来明细行（有效行需有借方或贷方金额）')
    if len(rows) > 50000:
        return err('明细超过5万行，请按月拆分导出后上传')

    with transaction.atomic():
        InternalBatch.objects.filter(business_unit=bu, year=year, month=month).delete()
        batch = InternalBatch.objects.create(
            business_unit=bu, year=year, month=month,
            filename=(f.name or '')[:200], row_count=len(rows),
            uploaded_by=getattr(getattr(request, 'pk_user', None), 'name', '') or '')
        for r in rows:
            r.batch = batch
        InternalEntry.objects.bulk_create(rows, batch_size=500)
    return ok({
        'batch': batch.to_dict(), 'rows': len(rows), 'skipped': skipped,
        'unmatched': [{'raw': k, 'count': v} for k, v in
                      sorted(unmatched.items(), key=lambda kv: -kv[1])[:50]],
    })


@csrf_exempt
@cw_required()
def internal_batches(request):
    """GET ?year=&month= → 各主体上传覆盖情况；DELETE 走 internal_batch_detail。"""
    denied = _page_denied(request, 'internal')
    if denied:
        return denied
    try:
        year = int(request.GET.get('year', ''))
        month = int(request.GET.get('month', ''))
    except Exception:
        return err('年份或月份无效')
    batches = {b.business_unit: b.to_dict()
               for b in InternalBatch.objects.filter(year=year, month=month)}
    return ok({'units': BUSINESS_UNITS,
               'batches': [batches.get(bu) for bu in BUSINESS_UNITS]})


@csrf_exempt
@cw_required()
def internal_batch_detail(request, bid):
    if request.method != 'DELETE':
        return err('方法不允许', 405)
    denied = _page_denied(request, 'internal')
    if denied:
        return denied
    if not (request.pk_role == 'super_admin' or _can_delete(request)):
        return err('无删除权限', 403)
    try:
        batch = InternalBatch.objects.get(id=bid)
    except InternalBatch.DoesNotExist:
        return err('批次不存在', 404)
    if not _can_access_bu(request, batch.business_unit):
        return err('无权操作该主体数据', 403)
    batch.delete()
    return ok({'deleted': bid})


def _net_positions(year, month):
    """{(bu, cp): Decimal 净头寸} — 仅已映射到有效主体的行。"""
    from django.db.models import Sum, F
    rows = (InternalEntry.objects
            .filter(year=year, month=month)
            .exclude(counterparty='')
            .values('business_unit', 'counterparty')
            .annotate(net=Sum(F('debit') - F('credit'))))
    return {(r['business_unit'], r['counterparty']): (r['net'] or Decimal('0')) for r in rows}


@csrf_exempt
@cw_required()
def internal_matrix(request):
    """GET ?year=&month= → 全集团差异矩阵 + KPI + 差异对 Top 列表 + 未识别统计。"""
    denied = _page_denied(request, 'internal')
    if denied:
        return denied
    try:
        year = int(request.GET.get('year', ''))
        month = int(request.GET.get('month', ''))
    except Exception:
        return err('年份或月份无效')
    pos = _net_positions(year, month)
    uploaded = set(InternalBatch.objects.filter(year=year, month=month)
                   .values_list('business_unit', flat=True))
    units = BUSINESS_UNITS
    cells = []          # 行=记账主体 A，列=对方 B：A账上对B净头寸
    pair_diffs = []     # 无序对 (A,B)：diff = pos(A,B) + pos(B,A)
    for a in units:
        row = []
        for b in units:
            row.append(float(pos.get((a, b), 0)) if a != b else None)
        cells.append(row)
    seen = set()
    for a in units:
        for b in units:
            if a == b or (b, a) in seen:
                continue
            seen.add((a, b))
            pa, pb = pos.get((a, b), Decimal('0')), pos.get((b, a), Decimal('0'))
            if not pa and not pb:
                continue
            both = a in uploaded and b in uploaded
            pair_diffs.append({
                'a': a, 'b': b,
                'a_net': float(pa), 'b_net': float(pb),
                'diff': float(pa + pb),
                'both_uploaded': both,
            })
    pair_diffs.sort(key=lambda p: -abs(p['diff']))
    total_ar = sum(float(v) for v in pos.values() if v > 0)
    total_ap = -sum(float(v) for v in pos.values() if v < 0)
    total_diff = sum(abs(p['diff']) for p in pair_diffs if p['both_uploaded'])
    matched_pairs = sum(1 for p in pair_diffs if p['both_uploaded'] and abs(p['diff']) < 0.005)
    unmatched_cnt = (InternalEntry.objects.filter(year=year, month=month, counterparty='')
                     .count())
    return ok({
        'units': units, 'uploaded': sorted(uploaded), 'cells': cells,
        'pairs': pair_diffs,
        'kpi': {
            'total_ar': total_ar, 'total_ap': total_ap,
            'total_diff': total_diff,
            'pairs_total': sum(1 for p in pair_diffs if p['both_uploaded']),
            'pairs_ok': matched_pairs,
            'unmatched_rows': unmatched_cnt,
        },
    })


def _auto_match(a_rows, b_rows):
    """镜像自动配对：|signed| 相等且方向相反的行 1:1 贪心配对（同额按日期先后）。
    返回 (a_out, b_out, matched_amount)，行 dict 带 match（配对组号）/None。"""
    a_out = [r.to_dict() for r in a_rows]
    b_out = [r.to_dict() for r in b_rows]
    from collections import defaultdict
    pool = defaultdict(list)   # key: 金额(取绝对值字符串) → b 侧待配行（按日期序）
    for i, r in enumerate(b_out):
        amt = abs(r['signed'])
        if amt:
            pool[f'{amt:.2f}'].append(i)
    gid = 0
    matched_amount = 0.0
    for r in a_out:
        r['match'] = None
    for r in b_out:
        r['match'] = None
    for ra in a_out:
        amt = abs(ra['signed'])
        if not amt:
            continue
        cand = pool.get(f'{amt:.2f}') or []
        # 镜像要求：signed 反号（A 记债权增，B 应记债务增 → B signed 为负）。
        # 同号候选（如双方都挂了应收）保留在池中，留给后续反号行。
        for k, j in enumerate(cand):
            if b_out[j]['signed'] * ra['signed'] < 0:
                cand.pop(k)
                gid += 1
                ra['match'] = gid
                b_out[j]['match'] = gid
                matched_amount += amt
                break
    return a_out, b_out, matched_amount


@csrf_exempt
@cw_required()
def internal_pair(request):
    """GET ?year=&month=&a=&b= → 两主体明细两栏比对 + 自动配对标记。"""
    denied = _page_denied(request, 'internal')
    if denied:
        return denied
    try:
        year = int(request.GET.get('year', ''))
        month = int(request.GET.get('month', ''))
    except Exception:
        return err('年份或月份无效')
    a = (request.GET.get('a') or '').strip()
    b = (request.GET.get('b') or '').strip()
    if a not in VALID_BUSINESS_UNITS or b not in VALID_BUSINESS_UNITS or a == b:
        return err('请选择两个不同的主体')
    a_rows = list(InternalEntry.objects
                  .filter(year=year, month=month, business_unit=a, counterparty=b)
                  .order_by('biz_date', 'id'))
    b_rows = list(InternalEntry.objects
                  .filter(year=year, month=month, business_unit=b, counterparty=a)
                  .order_by('biz_date', 'id'))
    a_out, b_out, matched_amount = _auto_match(a_rows, b_rows)
    a_net = sum(r['signed'] for r in a_out)
    b_net = sum(r['signed'] for r in b_out)
    uploaded = set(InternalBatch.objects.filter(year=year, month=month,
                                                business_unit__in=[a, b])
                   .values_list('business_unit', flat=True))
    return ok({
        'a': a, 'b': b,
        'a_rows': a_out, 'b_rows': b_out,
        'a_net': float(a_net), 'b_net': float(b_net),
        'diff': float(a_net + b_net),
        'matched_amount': matched_amount,
        'matched_pairs': sum(1 for r in a_out if r['match']),
        'a_uploaded': a in uploaded, 'b_uploaded': b in uploaded,
    })
