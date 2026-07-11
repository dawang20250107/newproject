"""Excel 导出样式 · 全系统单一真源。

品牌色 / 边框色 / 计算行底色 / 表头字体 / 金额数字格式此前在
report.py、paikuan、caiwu、各导入模板里各自内联，已出现细微不一致
（表头字号 10.5 vs 11、边框色 DDDDDD vs E3D5CC）。统一收敛到此处，
各导出点 import 使用，杜绝样式漂移。
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ── 设计令牌 ──────────────────────────────────────────────────────────────────
BRAND = 'C96342'            # 品牌主色（表头填充）
HEADER_TEXT = 'FFFFFF'      # 表头文字
BORDER_COLOR = 'E3D5CC'     # 细边框（暖灰，与品牌协调）
CALC_FILL_COLOR = 'FBEEE8'  # 计算行/合计行底色
HEADER_FONT_SIZE = 11
MONEY_FMT = '#,##0.00'      # 金额会计格式（千分位两位小数）
INT_FMT = '#,##0'           # 整数千分位


def header_fill():
    return PatternFill(fill_type='solid', fgColor=BRAND)


def header_font():
    return Font(bold=True, color=HEADER_TEXT, size=HEADER_FONT_SIZE)


def calc_fill():
    return PatternFill(fill_type='solid', fgColor=CALC_FILL_COLOR)


def thin_border():
    s = Side(style='thin', color=BORDER_COLOR)
    return Border(left=s, right=s, top=s, bottom=s)


def style_header_row(ws, ncols=None, height=22):
    """把第 1 行套成品牌表头（填充+白字+居中+行高）。"""
    n = ncols or ws.max_column
    fill, font = header_fill(), header_font()
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for col in range(1, n + 1):
        c = ws.cell(row=1, column=col)
        c.fill = fill
        c.font = font
        c.alignment = center
    ws.row_dimensions[1].height = height


def apply_money_format(ws, money_headers=()):
    """把表头名命中 money_headers 的列的数据单元格设为金额会计格式。返回命中的列号集合。"""
    if not money_headers:
        return set()
    wanted = set(money_headers)
    headers = [c.value for c in ws[1]]
    idxs = {i for i, h in enumerate(headers, 1) if h in wanted}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in idxs and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FMT
    return idxs


def append_total_row(ws, money_headers=(), label='合计', label_col=1):
    """在表末追加合计行：对 money_headers 命中列求和、label_col 写标签、整行套计算行底色。"""
    if ws.max_row < 2:
        return
    headers = [c.value for c in ws[1]]
    idxs = {i for i, h in enumerate(headers, 1) if h in set(money_headers)}
    if not idxs:
        return
    sums = {i: 0.0 for i in idxs}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            if cell.column in idxs and isinstance(cell.value, (int, float)):
                sums[cell.column] += float(cell.value)
    r = ws.max_row + 1
    fill = calc_fill()
    bold = Font(bold=True)
    ncols = len(headers)
    for col in range(1, ncols + 1):
        c = ws.cell(row=r, column=col)
        c.fill = fill
        c.font = bold
        if col == label_col and label_col not in idxs:
            c.value = label
        elif col in idxs:
            c.value = round(sums[col], 2)
            c.number_format = MONEY_FMT
    ws.row_dimensions[r].height = 20


def append_filter_snapshot(wb, pairs, title='筛选条件'):
    """在工作簿追加一页「筛选条件快照」：把导出时生效的筛选 (标签, 值) 列出，
    让导出件自证口径——收件人一眼看清这份数据是按什么条件筛出来的。
    pairs: [(label, value), ...]；value 为空的项自动跳过。"""
    rows = [(str(l), str(v)) for l, v in pairs if v not in (None, '', [], ())]
    ws = wb.create_sheet(title)
    ws.append(['筛选项', '取值'])
    style_header_row(ws, ncols=2)
    if not rows:
        ws.append(['（无）', '全部数据，未加筛选'])
    for l, v in rows:
        ws.append([l, v])
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 52
    ws.freeze_panes = 'A2'
